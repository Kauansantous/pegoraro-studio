from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
DOCX = OUT / "Ficha_do_Imovel_Pegoraro_Studio.docx"

BLACK = "11110F"
WARM = "68645C"
LIGHT = "DDD9D0"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.62)
sec.bottom_margin = Inches(0.62)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
styles["Normal"].font.size = Pt(10)
styles["Normal"].font.color.rgb = RGBColor.from_string(BLACK)

for name, size, face in [("Form Title", 28, "Georgia"), ("Form Section", 16, "Georgia"), ("Form Label", 9, "Arial")]:
    style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = face
    style._element.rPr.rFonts.set(qn("w:eastAsia"), face)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(BLACK)

def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right'):
        if edge in kwargs:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["val", "sz", "color", "space"]:
                if key in kwargs[edge]:
                    element.set(qn('w:{}'.format(key)), str(kwargs[edge][key]))

def footer():
    p = doc.sections[0].footer.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("PEGORARO STUDIO  |  FICHA DO IMÓVEL")
    r.font.name = "Arial"; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(WARM)

def heading(kicker, title, intro):
    p = doc.add_paragraph(style="Form Label")
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(kicker.upper())
    r.font.bold = True; r.font.color.rgb = RGBColor.from_string(WARM)
    p = doc.add_paragraph(style="Form Title")
    p.paragraph_format.space_after = Pt(8)
    p.add_run(title)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(intro)
    r.font.size = Pt(10); r.font.color.rgb = RGBColor.from_string(WARM)

def subheading(text):
    p = doc.add_paragraph(style="Form Section")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.add_run(text)

def field(label, lines=1, note=None):
    p = doc.add_paragraph(style="Form Label")
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(label.upper())
    r.font.bold = True; r.font.color.rgb = RGBColor.from_string(WARM)
    if note:
        r = p.add_run(f"  {note}")
        r.font.italic = True; r.font.bold = False
    for _ in range(lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run("________________________________________________________________________________")
        r.font.size = Pt(11); r.font.color.rgb = RGBColor.from_string(WARM)

def choices(label, opts):
    p = doc.add_paragraph(style="Form Label")
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run(label.upper() + "  ")
    r.font.bold = True; r.font.color.rgb = RGBColor.from_string(WARM)
    r = p.add_run("   ".join("[ ] " + x for x in opts))
    r.font.size = Pt(10); r.font.color.rgb = RGBColor.from_string(BLACK)

def new_page():
    if len(doc.paragraphs) > 0:
        doc.add_page_break()

footer()

# 1
heading("Briefing do anfitrião", "Ficha do imóvel", "Preencha à caneta tudo o que estiver definido. Onde ainda houver dúvida, escreva 'a confirmar'. Esta ficha reúne as informações necessárias para o Book Editorial, publicação e Guia do Hóspede.")
subheading("Antes de começar")
choices("Você já possui anúncio ativo", ["Airbnb", "Booking", "Outro", "Ainda não"])
field("Link do anúncio atual", 1)
field("Nome da pessoa que preenche esta ficha", 1)
field("Melhor WhatsApp e melhor e-mail para aprovações", 2)
subheading("Envie junto com esta ficha")
doc.add_paragraph("[ ] Link das fotos e vídeos existentes     [ ] Regras atuais da casa     [ ] Logotipo ou nome oficial do imóvel     [ ] Referências de anúncios que você gosta")
field("O que ainda falta enviar", 2)

# 2
new_page()
heading("02", "Identidade e posicionamento", "Informações que guiam o texto, a capa, a ordem das fotos e a percepção do anúncio.")
field("Nome oficial do imóvel", 1)
field("Endereço completo, complemento e ponto de referência", 2)
field("Cidade, bairro e localização que pode aparecer no anúncio", 1)
choices("Perfil de hóspede prioritário", ["Famílias", "Casais", "Grupos", "Corporativo", "Outro"])
field("Capacidade máxima de hóspedes", 1)
field("Três diferenciais que não podem faltar no anúncio", 3)
field("Como você quer que o imóvel seja percebido", 2)
field("Promessas, restrições ou informações que não podemos afirmar", 2)

# 3
new_page()
heading("03", "Estrutura e acomodações", "Dados técnicos que precisam estar corretos no anúncio e no manual.")
field("Quartos e suítes  tipo e quantidade de camas em cada quarto", 4)
field("Banheiros, lavabos e informação sobre água quente", 2)
field("Roupa de cama, toalhas, secador e itens extras  onde ficam", 2)
field("Ambientes disponíveis  sala, varanda, jardim, piscina, churrasqueira, praia e outros", 3)
field("Cozinha  cafeteira, fogão, forno, micro-ondas, água, freezer, lixo e reciclagem", 3)
field("Comodidades confirmadas  Wi-Fi, TV, streaming, ar-condicionado, máquina de lavar, garagem", 3)

# 4
new_page()
heading("04", "Chegada e acesso", "Detalhes que evitam dúvidas no check-in e diminuem chamadas de última hora.")
field("Horário de check-in e de check-out", 1)
field("Endereço para GPS e link de localização", 2)
field("Como o hóspede entra  portaria, portão, senha, chave ou recepção", 3)
field("Código, retirada de chave, local de chave reserva e controles", 3)
field("Estacionamento  vagas, onde parar, placas, restrições e instruções", 3)
field("Contato para problemas na chegada", 2)

# 5
new_page()
heading("05", "Funcionamento da casa", "Explique apenas o que o hóspede precisa saber para usar a casa com autonomia.")
field("Wi-Fi  nome da rede e senha", 1)
field("TV e streaming  controles, aplicativos e instruções", 2)
field("Ar-condicionado, ventiladores, aquecedores e controles", 2)
field("Lavanderia  se houver  máquina, varal e produtos", 2)
field("Piscina, hidro, sauna, churrasqueira, jardim ou praia  uso e horários", 3)
field("Itens ou áreas que não podem ser usados pelo hóspede", 2)

# 6
new_page()
heading("06", "Regras, segurança e saída", "Informações objetivas para manter a casa segura e a convivência tranquila.")
field("Silêncio e descanso  horários e orientação", 1)
field("Fumo, pets, visitas, festas e capacidade máxima", 3)
field("Cuidados do dia a dia e áreas restritas", 2)
field("Primeiros socorros, disjuntores, gás, extintor e instruções de segurança", 3)
field("Hospital ou UPA, SAMU, Bombeiros e responsável local  nome e telefone", 3)
field("Check-out  louça, lixo, luzes, ar-condicionado, portas, janelas, chave e aviso de saída", 4)

# 7
new_page()
heading("07", "Experiência do hóspede", "Recomendações e contatos que transformam uma boa estadia em uma experiência bem acompanhada.")
field("Nome do anfitrião e gestor de reservas", 1)
field("WhatsApp, e-mail, contato para emergência e link de avaliação", 3)
field("Restaurante, café ou padaria indicados  nome, endereço e link", 3)
field("Mercado, farmácia e conveniência  nome, endereço e link", 2)
field("Praia, passeio, trilha ou experiência imperdível", 2)
field("Café da manhã, pôr do sol, pausa ou ritual que você sugere", 2)
field("Mensagem de boas-vindas e assinatura do anfitrião", 2)

# 8
new_page()
heading("08", "Fotos, publicação e aprovação", "Última checagem para que a entrega seja publicada sem improvisos.")
field("Link das fotos e vídeos existentes", 1)
field("Ambientes ou detalhes obrigatórios para fotografar", 3)
field("Pessoas, objetos, áreas ou vistas que não podem aparecer", 2)
field("Foto preferida para capa e motivo", 2)
choices("Onde será publicado", ["Airbnb", "Booking", "Instagram", "Site próprio", "Outro"])
field("Comodidades que precisam ser conferidas antes de publicar", 2)
field("Responsável por aprovar o material  nome, WhatsApp e e-mail", 2)
field("Prazo ou data importante", 1)
choices("Autorização para usar o material no portfólio Pegoraro Studio", ["Sim", "Não", "A confirmar"])

doc.core_properties.title = "Ficha do Imóvel Pegoraro Studio"
doc.save(DOCX)

# Print-ready PDF counterpart. Kept deliberately simple so the host can write by hand.
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.pdfbase.pdfmetrics import stringWidth

PDF_DIR = OUT / "pdf"
PDF_DIR.mkdir(exist_ok=True)
PDF = PDF_DIR / "Ficha_do_Imovel_Pegoraro_Studio.pdf"
W, H = A4
c = canvas.Canvas(str(PDF), pagesize=A4)
ink = HexColor("#11110F")
warm = HexColor("#68645C")
line = HexColor("#B8B2A7")

pages = [
    ("BRIEFING DO ANFITRIÃO", "Ficha do imóvel", "Preencha à caneta tudo o que estiver definido. Onde ainda houver dúvida, escreva 'a confirmar'.", [
        ("O imóvel já tem anúncio ativo?  [ ] SIM  [ ] NÃO", 1), ("Se sim, qual é o link?", 1), ("Quem será nosso contato para aprovar o material?  Nome + WhatsApp + e-mail", 2), ("O que ainda será enviado?  [ ] FOTOS  [ ] VÍDEOS  [ ] REGRAS  [ ] OUTROS", 2)
    ]),
    ("02  IDENTIDADE E POSICIONAMENTO", "O imóvel em palavras", "Informações que guiam a capa, a ordem das fotos, o texto e a percepção do anúncio.", [
        ("Qual é o nome oficial do imóvel?", 1), ("Qual é o endereço completo e um ponto de referência?", 2), ("Em qual cidade/bairro fica? Quem você quer receber primeiro?", 2), ("Quantas pessoas podem se hospedar? Liste até 3 diferenciais.", 2), ("Complete: quero que o hóspede perceba esta casa como...", 2)
    ]),
    ("03  ESTRUTURA E ACOMODAÇÕES", "O que a casa oferece", "Dados técnicos para o anúncio e o manual do hóspede.", [
        ("Quantos quartos e camas existem? Ex.: Suíte 1 — 1 cama queen.", 3), ("Quantos banheiros/lavabos? Há água quente em todos?", 2), ("Quais ambientes o hóspede pode usar? [ ] SALA [ ] VARANDA [ ] JARDIM [ ] PISCINA [ ] OUTROS", 2), ("O que existe na cozinha? Marque os itens disponíveis e indique alguma instrução especial.", 3), ("Quais comodidades existem? [ ] WI-FI [ ] TV [ ] AR [ ] LAVANDERIA [ ] GARAGEM [ ] OUTRAS", 2)
    ]),
    ("04  CHEGADA E ACESSO", "Como o hóspede entra", "Detalhes para um check-in sem chamadas de última hora.", [
        ("Qual é o horário de entrada? E de saída?", 1), ("Qual endereço deve ser usado no GPS? Envie também o link do mapa.", 2), ("Como o hóspede entra na propriedade? Explique em até 3 passos.", 2), ("Existe senha, chave, controle ou chave reserva? Onde ficam?", 2), ("Há quantas vagas? Onde o hóspede deve estacionar?", 1), ("Quem deve ser chamado se houver problema na chegada?", 1)
    ]),
    ("05  FUNCIONAMENTO DA CASA", "Informações que evitam dúvidas", "Explique apenas o que o hóspede precisa saber para usar a casa com autonomia.", [
        ("Qual é o nome da rede Wi-Fi e a senha?", 1), ("Quais aparelhos precisam de instrução? Ex.: TV, ar, aquecedor.", 2), ("Existe algum cuidado para usar cozinha ou lavanderia?", 2), ("Há horário ou regra para piscina, hidro, sauna, churrasqueira ou jardim?", 2), ("Existe alguma área, armário ou equipamento que o hóspede não deve usar?", 1)
    ]),
    ("06  REGRAS SEGURANÇA E SAÍDA", "O que precisa ficar claro", "Informações objetivas para uma convivência tranquila e uma saída organizada.", [
        ("Quais são as regras para silêncio, fumo, pets, visitas e festas?", 2), ("Existe algum cuidado diário ou área restrita?", 1), ("Onde ficam extintor, disjuntores, registro de gás e primeiros socorros?", 2), ("Qual é o contato de emergência local? Informe hospital, UPA e responsável.", 2), ("O que o hóspede deve fazer antes de sair? Marque: [ ] LOUÇA [ ] LIXO [ ] LUZES [ ] AR [ ] PORTAS [ ] CHAVE", 2)
    ]),
    ("07  EXPERIÊNCIA DO HÓSPEDE", "O que vale recomendar", "Dicas e contatos que tornam a estadia mais fácil e mais especial.", [
        ("Quem é o anfitrião ou gestor? Informe nome, WhatsApp e e-mail.", 2), ("Indique 1 restaurante, 1 café, 1 mercado e 1 farmácia próximos.", 2), ("Qual passeio ou experiência você recomenda ao hóspede?", 2), ("Existe um momento especial da casa que vale sugerir?", 1), ("Há alguma mensagem de boas-vindas que você gostaria de deixar?", 2)
    ]),
    ("08  FOTOS PUBLICAÇÃO E APROVAÇÃO", "Antes de colocar no ar", "Última checagem para publicar com consistência e sem improvisos.", [
        ("Onde estão as fotos e vídeos? Cole aqui o link da pasta.", 1), ("Existe algum ambiente ou item que precisa aparecer? E algo que não pode aparecer?", 2), ("Você já tem uma foto de capa preferida? Qual e por quê?", 1), ("Onde o imóvel será publicado? [ ] AIRBNB [ ] BOOKING [ ] SITE [ ] OUTRO", 1), ("Quem aprova o material? Qual prazo? Podemos mostrar o projeto no portfólio? [ ] SIM [ ] NÃO", 2)
    ]),
]

for num, (kicker, title, intro, fields) in enumerate(pages, start=1):
    y = H - 54
    c.setFillColor(warm); c.setFont("Helvetica-Bold", 7.5); c.drawString(45, y, kicker)
    y -= 40
    c.setFillColor(ink); c.setFont("Times-Roman", 28)
    # lightweight wrap for title
    words = title.split(); current = ""; lines = []
    for word in words:
        trial = (current + " " + word).strip()
        if stringWidth(trial, "Times-Roman", 28) > 420:
            lines.append(current); current = word
        else: current = trial
    lines.append(current)
    for line_text in lines:
        c.drawString(45, y, line_text); y -= 30
    y -= 6
    c.setStrokeColor(line); c.line(45, y, W-45, y); y -= 20
    c.setFillColor(warm); c.setFont("Helvetica", 9.5)
    for chunk in [intro[i:i+96] for i in range(0, len(intro), 96)]:
        c.drawString(45, y, chunk); y -= 12
    y -= 9
    for label, rows in fields:
        c.setFillColor(warm); c.setFont("Helvetica-Bold", 7.2); c.drawString(45, y, label); y -= 13
        c.setStrokeColor(line)
        for _ in range(rows):
            c.line(45, y, W-45, y); y -= 23
        y -= 3
    c.setFillColor(warm); c.setFont("Helvetica", 7)
    c.drawString(45, 28, "PEGORARO STUDIO  |  FICHA DO IMÓVEL")
    c.drawRightString(W-45, 28, f"{num:02d}")
    c.showPage()
c.save()
print(DOCX)
print(PDF)
